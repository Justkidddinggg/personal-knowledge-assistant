from pathlib import Path

import pytest
from docx import Document

from app.ingestion.cleaning import clean_text
from app.ingestion.exceptions import EmptyDocumentError, UnsupportedFileTypeError
from app.ingestion.hashing import hash_file
from app.ingestion.loader import load_document


def test_clean_text_collapses_whitespace() -> None:
    assert clean_text("a   b\r\n\r\n\r\n\r\nc  ") == "a b\n\nc"


def test_load_txt(tmp_path: Path) -> None:
    f = tmp_path / "note.txt"
    f.write_text("สวัสดี hello world", encoding="utf-8")
    doc = load_document(f)
    assert doc.file_type == "txt"
    assert doc.pages[0].page is None
    assert "สวัสดี" in doc.pages[0].text


def test_load_markdown(tmp_path: Path) -> None:
    f = tmp_path / "readme.md"
    f.write_text("# Title\n\nSome content.", encoding="utf-8")
    assert load_document(f).file_type == "md"


def test_load_docx(tmp_path: Path) -> None:
    f = tmp_path / "doc.docx"
    d = Document()
    d.add_paragraph("First paragraph.")
    d.add_paragraph("Second paragraph.")
    d.save(f)
    doc = load_document(f)
    assert "First paragraph." in doc.pages[0].text


def test_unsupported_extension(tmp_path: Path) -> None:
    f = tmp_path / "image.png"
    f.write_bytes(b"fake")
    with pytest.raises(UnsupportedFileTypeError):
        load_document(f)


def test_empty_file_raises(tmp_path: Path) -> None:
    f = tmp_path / "empty.txt"
    f.write_text("   \n\n  ", encoding="utf-8")
    with pytest.raises(EmptyDocumentError):
        load_document(f)


def test_same_content_same_hash(tmp_path: Path) -> None:
    a, b = tmp_path / "a.txt", tmp_path / "b.txt"
    a.write_text("same", encoding="utf-8")
    b.write_text("same", encoding="utf-8")
    assert hash_file(a) == hash_file(b)


def test_pdf_has_page_numbers(tmp_path: Path) -> None:
    import fitz

    f = tmp_path / "sample.pdf"
    pdf = fitz.open()
    for i in range(2):
        page = pdf.new_page()
        page.insert_text((72, 72), f"Content of page {i + 1}")
    pdf.save(f)
    pdf.close()

    doc = load_document(f)
    assert [p.page for p in doc.pages] == [1, 2]
    assert "page 2" in doc.pages[1].text

def test_clean_text_fixes_thai_sara_am() -> None:
    assert clean_text("จ\u0e4d\u0e32กัด") == "จำกัด"


def test_clean_text_joins_hyphenated_line_breaks() -> None:
    assert clean_text("capabili-\nties and taxon-\nomy") == "capabilities and taxonomy"