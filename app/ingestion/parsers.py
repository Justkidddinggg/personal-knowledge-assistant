"""One parser per file format. Every parser returns list[ParsedPage]."""
import logging
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document

from app.ingestion.exceptions import CorruptDocumentError
from app.schemas.document import ParsedPage

logger = logging.getLogger(__name__)


def parse_pdf(path: Path) -> list[ParsedPage]:
    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise CorruptDocumentError(f"Cannot open PDF: {path.name}") from exc

    pages: list[ParsedPage] = []
    with doc:
        for index, page in enumerate(doc, start=1):  # human page numbers: 1-based
            pages.append(ParsedPage(text=page.get_text("text"), page=index))
    return pages


def parse_docx(path: Path) -> list[ParsedPage]:
    try:
        document = Document(path)
    except Exception as exc:
        raise CorruptDocumentError(f"Cannot open DOCX: {path.name}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # Tables are not in document.paragraphs, so read them separately.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [ParsedPage(text="\n".join(parts), page=None)]


def parse_text(path: Path) -> list[ParsedPage]:
    """Used for both .txt and .md (Markdown is kept as plain text)."""
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        logger.warning("%s is not valid UTF-8; retrying with replacement", path.name)
        text = path.read_text(encoding="utf-8", errors="replace")
    return [ParsedPage(text=text, page=None)]


PARSERS = {
    ".pdf": ("pdf", parse_pdf),
    ".docx": ("docx", parse_docx),
    ".txt": ("txt", parse_text),
    ".md": ("md", parse_text),
}